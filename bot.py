import os
import io
import json
import tempfile
import traceback
from datetime import datetime
from collections import defaultdict

import requests
from flask import Flask, request
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule

app = Flask(__name__)

BOT_TOKEN = os.environ.get("BOT_TOKEN", "").strip()
TELEGRAM_API = f"https://api.telegram.org/bot{BOT_TOKEN}"
TELEGRAM_FILE_API = f"https://api.telegram.org/file/bot{BOT_TOKEN}"

with open("config.json", "r", encoding="utf-8") as f:
    CONFIG = json.load(f)

CATEGORIES = CONFIG["categories"]
PLANS = CONFIG["plans"]
STORE_MAPPING = CONFIG["store_mapping"]
GROUP_MAPPING = CONFIG["group_mapping"]
DAYS_IN_MONTH = int(CONFIG.get("days_in_month", 30))

TITLE_FILL = "1F4E78"
HEADER_FILL = "D9EAF7"
STORE_FILL = "E2F0D9"


def tg(method, **kwargs):
    r = requests.post(f"{TELEGRAM_API}/{method}", timeout=60, **kwargs)
    r.raise_for_status()
    data = r.json()
    if not data.get("ok"):
        raise RuntimeError(data)
    return data["result"]


def send_message(chat_id, text):
    return tg("sendMessage", json={"chat_id": chat_id, "text": text})


def send_document(chat_id, path, caption=""):
    with open(path, "rb") as f:
        files = {"document": (os.path.basename(path), f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"chat_id": str(chat_id), "caption": caption}
        return tg("sendDocument", data=data, files=files)


def parse_iiko(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    period_text = str(ws["A3"].value or "")
    m = __import__("re").search(r"по\s+(\d{2})\.(\d{2})\.(\d{4})", period_text)
    if not m:
        raise ValueError("Не смогла определить конечную дату периода в строке A3.")
    day_end = int(m.group(1))
    month = int(m.group(2))
    year = int(m.group(3))

    blocks=[]
    enterprise=None
    subgroup=None
    items=[]

    for row in ws.iter_rows(min_row=4, values_only=True):
        a,b,c,d = row[:4]

        if a and str(a).strip() != "Итого":
            if items and enterprise and subgroup:
                blocks.append((enterprise, subgroup, items))
            enterprise = str(a).strip()
            subgroup = str(b).strip() if b else None
            items=[]

        elif enterprise and b and c is not None:
            if items and subgroup:
                blocks.append((enterprise, subgroup, items))
            subgroup = str(b).strip()
            items=[]

        elif enterprise and b and c is None:
            if items and subgroup:
                blocks.append((enterprise, subgroup, items))
                items=[]
            subgroup=None
            continue

        if enterprise and c is not None and d is not None:
            try:
                items.append((str(c).strip(), float(d)))
            except Exception:
                pass

    if items and enterprise and subgroup:
        blocks.append((enterprise, subgroup, items))

    return blocks, day_end, month, year


def analyze(path):
    blocks, day_end, month, year = parse_iiko(path)
    fact_by_store = {store: {c: 0.0 for c in CATEGORIES} for store in PLANS}
    unknown_stores=[]
    unknown_groups=[]

    for enterprise, subgroup, items in blocks:
        key = f"{enterprise}|||{subgroup}"
        store = STORE_MAPPING.get(key)
        if not store:
            unknown_stores.append(f"{enterprise} / {subgroup}")
            continue

        for raw_group, amount in items:
            cat = GROUP_MAPPING.get(raw_group)
            if not cat:
                unknown_groups.append(raw_group)
                continue
            fact_by_store[store][cat] += amount

    if unknown_stores:
        raise ValueError("Неизвестные магазины в новой выгрузке:\n" + "\n".join(sorted(set(unknown_stores))))
    if unknown_groups:
        raise ValueError("Новые группы товаров, которых нет в правилах:\n" + "\n".join(sorted(set(unknown_groups))))

    store_facts = {s: sum(v.values()) for s,v in fact_by_store.items()}
    total_fact = sum(store_facts.values())
    return fact_by_store, store_facts, total_fact, day_end, month, year


def style_title(ws, cell_range, text):
    ws.merge_cells(cell_range)
    c = ws[cell_range.split(":")[0]]
    c.value = text
    c.fill = PatternFill("solid", fgColor=TITLE_FILL)
    c.font = Font(bold=True, color="FFFFFF", size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")


def style_header(row_cells):
    for c in row_cells:
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def money(x):
    return float(x or 0)


def build_report(input_path, output_path):
    fact_by_store, store_facts, total_fact, days_fact, month, year = analyze(input_path)
    month_plan = sum(v["plan"] for v in PLANS.values())
    forecast_total = total_fact / days_fact * DAYS_IN_MONTH if days_fact else 0

    wb = Workbook()
    wb.remove(wb.active)

    ws_sum = wb.create_sheet("Сводка")
    ws_cat = wb.create_sheet("Категории")
    ws_store = wb.create_sheet("Магазины")
    ws_detail = wb.create_sheet("По магазинам и категориям")
    ws_all = wb.create_sheet("Все магазины")
    ws_rules = wb.create_sheet("Объединения")

    # Summary
    style_title(ws_sum, "A1:H1", f"Продажи SamalCakes — отчет за 1–{days_fact:02d}.{month:02d}.{year}")
    ws_sum.append([])
    ws_sum.append(["Показатель","Значение"])
    style_header(ws_sum[3])
    rows = [
        ("План месяца", month_plan),
        (f"Факт 1–{days_fact} сентября", total_fact),
        ("Дней факта", days_fact),
        ("Выполнение", total_fact/month_plan if month_plan else 0),
        ("Прогноз месяца", forecast_total),
        ("Прогноз, %", forecast_total/month_plan if month_plan else 0),
        ("Отклонение", forecast_total-month_plan),
        ("Дней осталось", DAYS_IN_MONTH-days_fact),
    ]
    for row in rows:
        ws_sum.append(row)
    for r in [4,5,8,10]:
        ws_sum[f"B{r}"].number_format = '#,##0'
    for r in [7,9]:
        ws_sum[f"B{r}"].number_format = '0.0%'

    # City summary
    city_agg = defaultdict(lambda: {"plan":0.0,"fact":0.0})
    for store in PLANS:
        city=store.split("/")[0]
        city_agg[city]["plan"] += PLANS[store]["plan"]
        city_agg[city]["fact"] += store_facts[store]

    style_title(ws_sum, "A14:H14", "Сводка по городам")
    headers=["Город","План, тг","Факт, тг","Выполнение","Прогноз, тг","Прогноз, %","Отклонение, тг","Доля факта"]
    for col, h in enumerate(headers,1):
        ws_sum.cell(15,col,h)
    style_header(ws_sum[15])

    rr=16
    for city, d in city_agg.items():
        p=d["plan"]; f=d["fact"]; fc=f/days_fact*DAYS_IN_MONTH if days_fact else 0
        vals=[city,p,f,f/p if p else 0,fc,fc/p if p else 0,fc-p,f/total_fact if total_fact else 0]
        for c,v in enumerate(vals,1):
            ws_sum.cell(rr,c,v)
        rr += 1

    # Categories
    style_title(ws_cat, "A1:H1", "Сводка по категориям")
    for c,h in enumerate(headers := ["Категория","План, тг","Факт, тг","Выполнение","Прогноз, тг","Прогноз, %","Отклонение, тг","Доля факта"],1):
        ws_cat.cell(3,c,h)
    style_header(ws_cat[3])
    cat_plan=defaultdict(float); cat_fact=defaultdict(float)
    for store,p in PLANS.items():
        for cat in CATEGORIES:
            cat_plan[cat] += p["categories"].get(cat,0)
            cat_fact[cat] += fact_by_store[store].get(cat,0)

    rr=4
    for cat in CATEGORIES:
        p=cat_plan[cat]; f=cat_fact[cat]; fc=f/days_fact*DAYS_IN_MONTH if days_fact else 0
        vals=[cat,p,f,f/p if p else 0,fc,fc/p if p else 0,fc-p,f/total_fact if total_fact else 0]
        for c,v in enumerate(vals,1):
            ws_cat.cell(rr,c,v)
        rr += 1

    # Stores
    style_title(ws_store, "A1:H1", "Прогноз по магазинам")
    store_headers=["Магазин","Город","План, тг","Факт, тг","Выполнение","Прогноз, тг","Прогноз, %","Отклонение, тг"]
    for c,h in enumerate(store_headers,1):
        ws_store.cell(3,c,h)
    style_header(ws_store[3])

    rr=4
    for store,pdata in PLANS.items():
        city=store.split("/")[0]; p=pdata["plan"]; f=store_facts[store]; fc=f/days_fact*DAYS_IN_MONTH if days_fact else 0
        vals=[store,city,p,f,f/p if p else 0,fc,fc/p if p else 0,fc-p]
        for c,v in enumerate(vals,1):
            ws_store.cell(rr,c,v)
        rr += 1

    # All stores with numbering
    style_title(ws_all, "A1:I1", "Все магазины — контрольный список")
    all_headers=["№","Город","Магазин","План, тг","Факт, тг","Выполнение","Прогноз, тг","Прогноз, %","Отклонение, тг"]
    for c,h in enumerate(all_headers,1):
        ws_all.cell(3,c,h)
    style_header(ws_all[3])
    rr=4
    city_counts=defaultdict(int)
    for n,(store,pdata) in enumerate(PLANS.items(),1):
        city=store.split("/")[0]; city_counts[city]+=1
        p=pdata["plan"]; f=store_facts[store]; fc=f/days_fact*DAYS_IN_MONTH if days_fact else 0
        vals=[n,city,store,p,f,f/p if p else 0,fc,fc/p if p else 0,fc-p]
        for c,v in enumerate(vals,1):
            ws_all.cell(rr,c,v)
        rr += 1
    ws_all["K1"]="Всего магазинов"; ws_all["L1"]=len(PLANS)
    kr=2
    for city,count in city_counts.items():
        ws_all.cell(kr,11,city); ws_all.cell(kr,12,count); kr+=1

    # Detail
    style_title(ws_detail, "A1:H1", "По каждому магазину — план, факт и категории")
    ws_detail["A2"]="Период факта"; ws_detail["B2"]=f"01.{month:02d}.{year}–{days_fact:02d}.{month:02d}.{year}"
    ws_detail["A3"]="Дней факта"; ws_detail["B3"]=days_fact
    ws_detail["A4"]="Дней в месяце"; ws_detail["B4"]=DAYS_IN_MONTH
    detail_headers=["Магазин / категория","План","Факт","Выполнение","Прогноз","Прогноз, %","Отставание","Доля в магазине, %"]
    for c,h in enumerate(detail_headers,1):
        ws_detail.cell(6,c,h)
    style_header(ws_detail[6])
    rr=7
    for store,pdata in PLANS.items():
        p=pdata["plan"]; f=store_facts[store]; fc=f/days_fact*DAYS_IN_MONTH if days_fact else 0
        vals=[store,p,f,f/p if p else 0,fc,fc/p if p else 0,fc-p,1]
        for c,v in enumerate(vals,1):
            ws_detail.cell(rr,c,v)
            ws_detail.cell(rr,c).fill=PatternFill("solid",fgColor=STORE_FILL)
            ws_detail.cell(rr,c).font=Font(bold=True)
        rr += 1
        for cat in CATEGORIES:
            cp=pdata["categories"].get(cat,0); cf=fact_by_store[store].get(cat,0); cfc=cf/days_fact*DAYS_IN_MONTH if days_fact else 0
            vals=[cat,cp,cf,cf/cp if cp else 0,cfc,cfc/cp if cp else 0,cfc-cp,cf/f if f else 0]
            for c,v in enumerate(vals,1):
                ws_detail.cell(rr,c,v)
            rr += 1
        rr += 1

    # Rules
    style_title(ws_rules, "A1:C1", "Правила объединения категорий")
    for c,h in enumerate(["Исходная группа iiko","Категория отчета","Примечание"],1):
        ws_rules.cell(3,c,h)
    style_header(ws_rules[3])
    rr=4
    for raw,cat in GROUP_MAPPING.items():
        ws_rules.cell(rr,1,raw); ws_rules.cell(rr,2,cat); rr+=1

    # Formatting
    for ws in [ws_sum, ws_cat, ws_store, ws_all, ws_detail]:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value,(int,float)):
                    if cell.column in [4,6,8] and ws.title in ["Категории","Сводка","По магазинам и категориям"]:
                        pass

    # Sheet-specific formats
    for ws, start_row, end_row, money_cols, pct_cols in [
        (ws_sum,16,15+len(city_agg),[2,3,5,7],[4,6,8]),
        (ws_cat,4,3+len(CATEGORIES),[2,3,5,7],[4,6,8]),
        (ws_store,4,3+len(PLANS),[3,4,6,8],[5,7]),
        (ws_all,4,3+len(PLANS),[4,5,7,9],[6,8]),
    ]:
        for r in range(start_row,end_row+1):
            for c in money_cols: ws.cell(r,c).number_format='#,##0'
            for c in pct_cols: ws.cell(r,c).number_format='0.0%'

    for r in range(7, ws_detail.max_row+1):
        for c in [2,3,5,7]: ws_detail.cell(r,c).number_format='#,##0'
        for c in [4,6,8]: ws_detail.cell(r,c).number_format='0.0%'

    # Widths and freeze panes
    for ws in [ws_sum,ws_cat,ws_store,ws_all,ws_detail,ws_rules]:
        ws.freeze_panes = "A4" if ws.title in ["Магазины","Все магазины"] else ("A7" if ws.title=="По магазинам и категориям" else None)
        for col in range(1, ws.max_column+1):
            letter=__import__("openpyxl").utils.get_column_letter(col)
            max_len=0
            for cell in ws[letter]:
                if cell.value is not None:
                    max_len=max(max_len,len(str(cell.value)))
            ws.column_dimensions[letter].width=min(max(max_len+2,12),36)

    # Conditional scales
    for ws, rng in [
        (ws_sum, f"F16:F{15+len(city_agg)}"),
        (ws_cat, f"F4:F{3+len(CATEGORIES)}"),
        (ws_store, f"G4:G{3+len(PLANS)}"),
        (ws_all, f"H4:H{3+len(PLANS)}"),
        (ws_detail, f"F7:F{ws_detail.max_row}"),
    ]:
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type='min', start_color='F4CCCC',
            mid_type='percentile', mid_value=50, mid_color='FFF2CC',
            end_type='max', end_color='D9EAD3'
        ))

    wb.save(output_path)
    return total_fact, month_plan, forecast_total, days_fact


@app.get("/")
def home():
    return "Samal Report Bot работает ✅"


@app.get("/setup")
def setup_webhook():
    if not BOT_TOKEN:
        return "BOT_TOKEN не задан", 500
    external = os.environ.get("RENDER_EXTERNAL_URL","").rstrip("/")
    if not external:
        return "RENDER_EXTERNAL_URL пока недоступен", 500
    url = external + "/telegram"
    result = tg("setWebhook", json={"url": url, "drop_pending_updates": True})
    return f"Webhook установлен: {url} -> {result}"


@app.post("/telegram")
def telegram_webhook():
    update = request.get_json(silent=True) or {}
    try:
        msg = update.get("message") or {}
        chat = msg.get("chat") or {}
        chat_id = chat.get("id")
        if not chat_id:
            return "ok"

        text = msg.get("text","")
        if text == "/start":
            send_message(chat_id,
                "Привет! 👋\n"
                "Отправь мне свежий Excel-файл (.xlsx) из iiko.\n"
                "Я сделаю отчет: сводка, города, все магазины, категории и прогноз."
            )
            return "ok"

        doc = msg.get("document")
        if not doc:
            send_message(chat_id, "Отправь, пожалуйста, Excel-файл .xlsx из iiko.")
            return "ok"

        name = doc.get("file_name","")
        if not name.lower().endswith(".xlsx"):
            send_message(chat_id, "Мне нужен файл Excel с расширением .xlsx.")
            return "ok"

        send_message(chat_id, "Файл получила ✅ Считаю отчет...")

        f_info = tg("getFile", json={"file_id": doc["file_id"]})
        file_path = f_info["file_path"]
        content = requests.get(f"{TELEGRAM_FILE_API}/{file_path}", timeout=60).content

        with tempfile.TemporaryDirectory() as td:
            inp = os.path.join(td, "iiko.xlsx")
            out = os.path.join(td, "Samal_Report.xlsx")
            with open(inp,"wb") as f:
                f.write(content)

            total_fact, plan, forecast, days = build_report(inp, out)
            pct = forecast/plan if plan else 0

            caption = (
                f"Готово ✅\n"
                f"Факт за 1–{days}: {total_fact:,.0f} тг\n"
                f"Прогноз: {forecast:,.0f} тг ({pct:.1%})"
            ).replace(",", " ")

            send_document(chat_id, out, caption=caption)

    except Exception as e:
        traceback.print_exc()
        try:
            chat_id = ((update.get("message") or {}).get("chat") or {}).get("id")
            if chat_id:
                send_message(chat_id, "Не получилось обработать файл.\n\n" + str(e)[:3000])
        except Exception:
            pass

    return "ok"
